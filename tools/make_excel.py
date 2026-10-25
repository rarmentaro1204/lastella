#!/usr/bin/env python3
"""Genera i file Excel per i clienti di La Stella Catering.

  python3 tools/make_excel.py

Produce:
  template/Invitati_e_Tavoli_La_Stella.xlsx   (modello vuoto da consegnare agli sposi)
  esempio/Esempio_compilato.xlsx              (esempio con ~115 invitati finti, per test/demo)

Foglio 1  "Invitati"  - anagrafica + allergie / intolleranze / dieta
Foglio 2  "Tavoli"    - forma del tavolo (tondo / rettangolare) + invitato per ogni posto
Foglio 3  "Leggimi"   - istruzioni, legenda, riepilogo automatico
"""
import os
import random
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# --- palette La Stella ------------------------------------------------------
NAVY, SAND, PAPER, INK = "2E3540", "A8907A", "F4F1EC", "23272E"
TERRA, AMBER, SAGE, GREY = "C46A4F", "D9A441", "8FA383", "E6E2DB"

ALLERGIE = [
    ("glutine", "Glutine (cereali)"), ("crostacei", "Crostacei"), ("uova", "Uova"),
    ("pesce", "Pesce"), ("arachidi", "Arachidi"), ("soia", "Soia"),
    ("latte", "Latte e derivati"), ("guscio", "Frutta a guscio"), ("sedano", "Sedano"),
    ("senape", "Senape"), ("sesamo", "Sesamo"), ("solfiti", "Solfiti"),
    ("lupini", "Lupini"), ("molluschi", "Molluschi"),
]
INTOLLERANZE = [
    ("celiachia", "Celiachia"), ("lattosio", "Lattosio"), ("nichel", "Nichel"),
    ("istamina", "Istamina"), ("fruttosio", "Fruttosio"), ("favismo", "Favismo"),
]
DIETE = ["Vegetariano", "Vegano", "Fruttariano", "Pescetariano", "Crudista", "Halal", "Kosher"]
FASCE = ["Adulto", "Bambino", "Neonato"]

N_GUEST_ROWS = 300      # righe disponibili in "Invitati"
N_TABLES = 40           # tavoli disponibili in "Tavoli"
MAX_SEATS = 24          # posti massimi per tavolo

thin = Side(style="thin", color="D5CFC5")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def hdr(cell, bg, fg="FFFFFF", rotate=False):
    cell.font = Font(name="Calibri", bold=True, color=fg, size=10)
    cell.fill = fill(bg)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True,
                               text_rotation=90 if rotate else 0)
    cell.border = BORDER


def build(path, sample=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Invitati"
    wt = wb.create_sheet("Tavoli")
    wr = wb.create_sheet("Leggimi")

    # ---------------------------------------------------------------- Invitati
    cols = ["Nome", "Cognome", "Fascia", "Dieta"]
    cols += [f"Allergia – {lab}" for _, lab in ALLERGIE]
    cols += [f"Intolleranza – {lab}" for _, lab in INTOLLERANZE]
    cols += ["Altre allergie / intolleranze (scrivere)", "Note per la cucina", "Nome completo", "Stato tavolo"]
    c_first_all = 5
    c_first_int = c_first_all + len(ALLERGIE)
    c_other = c_first_int + len(INTOLLERANZE)
    c_note = c_other + 1
    c_full = c_note + 1
    c_stato = c_full + 1
    last = N_GUEST_ROWS + 1

    for i, name in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=name)
        if i <= 3:
            hdr(c, NAVY)
        elif i == 4:
            hdr(c, SAGE, "FFFFFF")
        elif i < c_first_int:
            hdr(c, TERRA, rotate=True)
        elif i < c_other:
            hdr(c, AMBER, INK, rotate=True)
        elif i in (c_other, c_note):
            hdr(c, NAVY)
        else:
            hdr(c, SAND)
    ws.row_dimensions[1].height = 130

    widths = {1: 16, 2: 18, 3: 11, 4: 15, c_other: 28, c_note: 28, c_full: 24, c_stato: 14}
    for i in range(1, len(cols) + 1):
        ws.column_dimensions[L(i)].width = widths.get(i, 4.6)
    ws.freeze_panes = "C2"

    dv_fascia = DataValidation(type="list", formula1='"' + ",".join(FASCE) + '"', allow_blank=True)
    dv_dieta = DataValidation(type="list", formula1='"' + ",".join(DIETE) + '"', allow_blank=True)
    dv_x = DataValidation(type="list", formula1='"X"', allow_blank=True, showErrorMessage=False)
    for dv in (dv_fascia, dv_dieta, dv_x):
        ws.add_data_validation(dv)
    dv_fascia.add(f"C2:C{last}")
    dv_dieta.add(f"D2:D{last}")
    dv_x.add(f"{L(c_first_all)}2:{L(c_other - 1)}{last}")

    for r in range(2, last + 1):
        for ci in range(1, len(cols) + 1):
            ws.cell(row=r, column=ci).border = BORDER
        for ci in range(c_first_all, c_other):
            ws.cell(row=r, column=ci).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=c_full, value=f'=IF(OR(A{r}<>"",B{r}<>""),TRIM(A{r}&" "&B{r}),"")')
        t_rng = f"Tavoli!$E$2:${L(4 + MAX_SEATS)}${N_TABLES + 1}"
        t_num = f"Tavoli!$A$2:$A${N_TABLES + 1}"
        full = f"{L(c_full)}{r}"
        ws.cell(row=r, column=c_stato, value=(
            f'=IF({full}="","",IF(COUNTIF({t_rng},{full})=0,"Da assegnare",'
            f'IF(COUNTIF({t_rng},{full})>1,"DUPLICATO","Tavolo "&SUMPRODUCT(({t_rng}={full})*{t_num}))))'))
        for ci in (c_full, c_stato):
            ws.cell(row=r, column=ci).fill = fill(GREY)
            ws.cell(row=r, column=ci).font = Font(color="6B6B6B", size=9)

    ws.conditional_formatting.add(
        f"{L(c_first_all)}2:{L(c_first_int - 1)}{last}",
        CellIsRule(operator="equal", formula=['"X"'], fill=fill(TERRA), font=Font(color="FFFFFF", bold=True)))
    ws.conditional_formatting.add(
        f"{L(c_first_int)}2:{L(c_other - 1)}{last}",
        CellIsRule(operator="equal", formula=['"X"'], fill=fill(AMBER), font=Font(color=INK, bold=True)))
    ws.conditional_formatting.add(
        f"{L(c_stato)}2:{L(c_stato)}{last}",
        FormulaRule(formula=[f'{L(c_stato)}2="DUPLICATO"'], fill=fill("F2B8A8"), font=Font(bold=True, color="8A2A12")))
    ws.conditional_formatting.add(
        f"{L(c_stato)}2:{L(c_stato)}{last}",
        FormulaRule(formula=[f'{L(c_stato)}2="Da assegnare"'], font=Font(color="B07A1A", italic=True)))

    # ------------------------------------------------------------------ Tavoli
    th = ["N. tavolo", "Nome tavolo (facoltativo)", "Forma tavolo", "N. posti"]
    th += [f"Posto {i}" for i in range(1, MAX_SEATS + 1)]
    th += ["Occupati", "Controllo"]
    for i, name in enumerate(th, start=1):
        c = wt.cell(row=1, column=i, value=name)
        hdr(c, NAVY if i <= 4 else (SAND if i <= 4 + MAX_SEATS else SAGE))
    wt.row_dimensions[1].height = 34
    wt.column_dimensions["A"].width = 9
    wt.column_dimensions["B"].width = 22
    wt.column_dimensions["C"].width = 15
    wt.column_dimensions["D"].width = 9
    for i in range(5, 5 + MAX_SEATS):
        wt.column_dimensions[L(i)].width = 19
    wt.column_dimensions[L(5 + MAX_SEATS)].width = 10
    wt.column_dimensions[L(6 + MAX_SEATS)].width = 22
    wt.freeze_panes = "E2"

    tl = N_TABLES + 1
    s1, sN = "E", L(4 + MAX_SEATS)
    dv_forma = DataValidation(type="list", formula1='"Tondo,Rettangolare"', allow_blank=True,
                              promptTitle="Forma", prompt="Scegli tavolo Tondo o Rettangolare (imperiale)", showInputMessage=True)
    dv_posti = DataValidation(type="whole", operator="between", formula1="1", formula2=str(MAX_SEATS),
                              errorTitle="N. posti", error=f"Inserisci un numero da 1 a {MAX_SEATS}", allow_blank=True)
    dv_seat = DataValidation(type="list", formula1=f"=Invitati!${L(c_full)}$2:${L(c_full)}${last}", allow_blank=True,
                             errorTitle="Invitato non trovato",
                             error="Scegli un nome dall'elenco (deve prima essere inserito nel foglio Invitati).",
                             showErrorMessage=True)
    for dv in (dv_forma, dv_posti, dv_seat):
        wt.add_data_validation(dv)
    dv_forma.add(f"C2:C{tl}")
    dv_posti.add(f"D2:D{tl}")
    dv_seat.add(f"{s1}2:{sN}{tl}")

    for r in range(2, tl + 1):
        wt.cell(row=r, column=1, value=r - 1)
        for ci in range(1, len(th) + 1):
            wt.cell(row=r, column=ci).border = BORDER
        wt.cell(row=r, column=1).alignment = Alignment(horizontal="center")
        wt.cell(row=r, column=1).font = Font(bold=True)
        wt.cell(row=r, column=4).alignment = Alignment(horizontal="center")
        occ = L(5 + MAX_SEATS)
        wt.cell(row=r, column=5 + MAX_SEATS, value=f"=COUNTA({s1}{r}:{sN}{r})")
        wt.cell(row=r, column=6 + MAX_SEATS, value=(
            f'=IF(AND(C{r}="",D{r}="",{occ}{r}=0),"",IF(OR(C{r}="",D{r}=""),"Indicare forma e posti",'
            f'IF({occ}{r}>D{r},"Troppi invitati",IF({occ}{r}<D{r},"Posti liberi: "&(D{r}-{occ}{r}),"OK"))))'))
        wt.cell(row=r, column=5 + MAX_SEATS).alignment = Alignment(horizontal="center")
    # posti oltre "N. posti" -> grigi ; nomi doppi -> rosso
    wt.conditional_formatting.add(
        f"{s1}2:{sN}{tl}",
        FormulaRule(formula=[f'AND({s1}2<>"",COUNTIF(${s1}$2:${sN}${tl},{s1}2)>1)'],
                    fill=fill("F2B8A8"), font=Font(bold=True, color="8A2A12")))
    wt.conditional_formatting.add(
        f"{s1}2:{sN}{tl}",
        FormulaRule(formula=[f'AND($D2<>"",COLUMN()-4>$D2)'], fill=fill("D9D5CE")))
    wt.conditional_formatting.add(
        f"{L(6 + MAX_SEATS)}2:{L(6 + MAX_SEATS)}{tl}",
        FormulaRule(formula=[f'{L(6 + MAX_SEATS)}2="OK"'], fill=fill("DCE8D5"), font=Font(color="3E5A32", bold=True)))
    wt.conditional_formatting.add(
        f"{L(6 + MAX_SEATS)}2:{L(6 + MAX_SEATS)}{tl}",
        FormulaRule(formula=[f'ISNUMBER(SEARCH("Troppi",{L(6 + MAX_SEATS)}2))'], fill=fill("F2B8A8"), font=Font(color="8A2A12", bold=True)))

    # ----------------------------------------------------------------- Leggimi
    wr.sheet_view.showGridLines = False
    wr.column_dimensions["A"].width = 3
    wr.column_dimensions["B"].width = 34
    wr.column_dimensions["C"].width = 90
    r = 2
    wr.cell(row=r, column=2, value="L A   S T E L L A").font = Font(name="Georgia", size=20, color=NAVY, bold=True)
    r += 1
    wr.cell(row=r, column=2, value="N E W   C O N C E P T   B A N Q U E T I N G").font = Font(size=9, color=SAND)
    r += 2
    wr.cell(row=r, column=2, value="Elenco invitati e disposizione ai tavoli").font = Font(size=14, bold=True, color=INK)
    r += 2

    def row(label, text, bold=True, bg=None):
        nonlocal r
        a = wr.cell(row=r, column=2, value=label)
        b = wr.cell(row=r, column=3, value=text)
        a.font = Font(bold=bold, color=INK)
        a.alignment = Alignment(vertical="top", wrap_text=True)
        b.alignment = Alignment(vertical="top", wrap_text=True)
        if bg:
            a.fill = fill(bg); b.fill = fill(bg)
        wr.row_dimensions[r].height = max(18, 15 * (1 + len(text) // 95))
        r += 1

    row("Come si compila", "Compilate i due fogli in ordine: prima «Invitati», poi «Tavoli». Poi rimandateci il file così com'è (non cambiate nomi dei fogli o l'ordine delle colonne).")
    row("1 · Foglio «Invitati»", "Una riga per persona: Nome, Cognome, Fascia (Adulto / Bambino / Neonato) e, se serve, Dieta. Mettete una X nelle colonne delle ALLERGIE (rosso) e delle INTOLLERANZE (giallo) che riguardano quella persona. Per tutto il resto usate le due colonne di testo a destra.")
    row("   Fascia", "Adulto · Bambino (circa 3-10 anni) · Neonato (0-2 anni). Bambini e neonati occupano comunque un posto a tavola.")
    row("   Dieta", "Vegetariano · Vegano · Fruttariano · Pescetariano · Crudista · Halal · Kosher. Lasciate vuoto se nessuna.")
    row("   Allergie", "Sono i 14 allergeni a dichiarazione obbligatoria (Reg. UE 1169/2011). Una allergia è un rischio per la salute: segnalatela sempre, anche se lieve.")
    row("   Intolleranze", "Celiachia, lattosio, nichel, istamina, fruttosio, favismo. Se manca la vostra, scrivetela in «Altre allergie / intolleranze».")
    row("2 · Foglio «Tavoli»", "Per ogni tavolo: scegliete la FORMA (Tondo o Rettangolare/imperiale), indicate il NUMERO DI POSTI e poi, nelle colonne «Posto 1…», scegliete dal menu a tendina il nome dell'invitato (l'elenco si alimenta da «Invitati»).")
    row("   Tavoli tondi", "Da un minimo di 8 a un massimo di 11 persone per tavolo (bambini e neonati compresi). I posti sono numerati in senso orario partendo dal posto 1.")
    row("   Tavoli rettangolari", "I posti si dividono sui due lati lunghi: dal posto 1 si riempie il lato superiore da sinistra a destra (metà posti, +1 se dispari), poi il lato inferiore da sinistra a destra.")
    row("   Controlli automatici", "Colonna «Controllo»: segnala posti liberi / troppi invitati. Un nome ripetuto in due posti diventa rosso. In «Invitati» l'ultima colonna dice se la persona è già assegnata a un tavolo.")
    row("Scadenze", "Le scelte di allestimento vanno confermate entro la data indicata dal vostro referente. Elenco invitati e tavoli: inviateci una prima versione appena possibile, potrete aggiornarla fino a pochi giorni prima dell'evento.")
    r += 1
    wr.cell(row=r, column=2, value="Riepilogo automatico").font = Font(size=12, bold=True, color=INK)
    r += 1
    iv = f"Invitati!$C$2:$C${last}"
    rows = [
        ("Invitati inseriti", f'=SUMPRODUCT(--(Invitati!$A$2:$A${last}&Invitati!$B$2:$B${last}<>""))'),
        ("   di cui adulti", f'=COUNTIF({iv},"Adulto")'),
        ("   di cui bambini", f'=COUNTIF({iv},"Bambino")'),
        ("   di cui neonati", f'=COUNTIF({iv},"Neonato")'),
        ("Tavoli compilati", f'=COUNTIF(Tavoli!$D$2:$D${tl},">0")'),
        ("Posti totali", f"=SUM(Tavoli!$D$2:$D${tl})"),
        ("Invitati ancora da assegnare", f'=COUNTIF(Invitati!${L(c_stato)}$2:${L(c_stato)}${last},"Da assegnare")'),
        ("Segnalazioni di allergia (totale X)", f'=COUNTIF(Invitati!${L(c_first_all)}$2:${L(c_first_int - 1)}${last},"X")'),
    ]
    for lab, f in rows:
        wr.cell(row=r, column=2, value=lab).font = Font(color=INK)
        c = wr.cell(row=r, column=3, value=f)
        c.alignment = Alignment(horizontal="left")
        c.font = Font(bold=True, color=NAVY)
        r += 1
    r += 1
    wr.cell(row=r, column=2, value="Per qualsiasi dubbio: info@cateringlastella.com · +39 339 385 8184").font = Font(italic=True, color=SAND)

    for sh, color in ((ws, NAVY), (wt, SAND), (wr, SAGE)):
        sh.sheet_properties.tabColor = color

    # ---------------------------------------------------------------- esempio
    if sample:
        fill_sample(ws, wt, c_first_all, c_first_int, c_other, c_note)

    wb.save(path)
    print("scritto", path)


def fill_sample(ws, wt, c_all, c_int, c_other, c_note):
    rnd = random.Random(2027)
    nomi_m = "Marco Luca Andrea Matteo Davide Simone Paolo Stefano Giorgio Fabio Enrico Alessio Gabriele Nicola Daniele Riccardo".split()
    nomi_f = "Giulia Chiara Elena Sara Martina Valentina Federica Silvia Laura Paola Alice Elisa Francesca Beatrice Marta Irene".split()
    cogn = "Rossi Bianchi Ferrari Russo Romano Colombo Ricci Marino Greco Bruno Gallo Conti De Luca Mancini Costa Giordano Rizzo Lombardi Moretti Barbieri Fontana Santoro Mariani Rinaldi Caruso Ferrara Galli Martini Leone Longo Gentile Martinelli Vitale Lombardo Serra Coppola".split()
    people = []
    seen = set()
    def uniq():
        while True:
            n = rnd.choice(nomi_m + nomi_f)
            c = rnd.choice(cogn)
            if (n, c) not in seen:
                seen.add((n, c)); return n, c
    for i in range(104):
        people.append(uniq() + ("Adulto",))
    for i in range(9):
        people.append(uniq() + ("Bambino",))
    for i in range(2):
        people.append(uniq() + ("Neonato",))
    rnd.shuffle(people)
    # esigenze speciali
    allergie_keys = [k for k, _ in ALLERGIE]
    int_keys = [k for k, _ in INTOLLERANZE]
    for idx, (n, c, f) in enumerate(people):
        r = idx + 2
        ws.cell(row=r, column=1, value=n)
        ws.cell(row=r, column=2, value=c)
        ws.cell(row=r, column=3, value=f)
        x = rnd.random()
        if x < 0.07:
            ws.cell(row=r, column=4, value="Vegetariano")
        elif x < 0.10:
            ws.cell(row=r, column=4, value="Vegano")
        elif x < 0.11:
            ws.cell(row=r, column=4, value="Pescetariano")
        elif x < 0.12:
            ws.cell(row=r, column=4, value="Fruttariano")
        if rnd.random() < 0.14:
            for k in rnd.sample(allergie_keys, rnd.choice([1, 1, 2])):
                ws.cell(row=r, column=c_all + allergie_keys.index(k), value="X")
        if rnd.random() < 0.12:
            for k in rnd.sample(int_keys, 1):
                ws.cell(row=r, column=c_int + int_keys.index(k), value="X")
        if idx == 5:
            ws.cell(row=r, column=c_other, value="Allergia al kiwi")
        if idx == 9:
            ws.cell(row=r, column=c_note, value="Seggiolone per favore")
    # tavoli: 11 tondi (8-11 posti) + 1 rettangolare per il resto
    names = [f"{n} {c}" for n, c, _ in people]
    pos = 0
    forms = []
    remaining = len(names)
    t = 1
    while remaining > 0:
        if t == 12:
            seats, forma = remaining, "Rettangolare"
        else:
            seats = min(remaining, rnd.choice([9, 10, 10, 10, 11]))
            forma = "Tondo"
        forms.append((t, forma, seats))
        remaining -= seats
        t += 1
    for (t, forma, seats) in forms:
        r = t + 1
        wt.cell(row=r, column=2, value=f"Tavolo {t}" if t < 12 else "Imperiale")
        wt.cell(row=r, column=3, value=forma)
        wt.cell(row=r, column=4, value=seats)
        for s in range(seats):
            wt.cell(row=r, column=5 + s, value=names[pos]); pos += 1
    wt.cell(row=2, column=2, value="Sposi")


if __name__ == "__main__":
    os.makedirs(os.path.join(ROOT, "template"), exist_ok=True)
    os.makedirs(os.path.join(ROOT, "esempio"), exist_ok=True)
    build(os.path.join(ROOT, "template", "Invitati_e_Tavoli_La_Stella.xlsx"))
    build(os.path.join(ROOT, "esempio", "Esempio_compilato.xlsx"), sample=True)
